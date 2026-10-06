import io

import pytest
from docx import Document
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from masking.app import app
from masking.detectors import find_pii, inn_ok, luhn_ok, snils_ok
from masking.docx_proc import mask_docx
from masking.masker import Masker
from masking.xlsx_proc import mask_xlsx

client = TestClient(app)


def kinds(text):
    return [(text[m.start:m.end], m.kind) for m in find_pii(text)]


def test_checksums():
    assert inn_ok("500100732259") and not inn_ok("500100732258")
    assert inn_ok("7707083893") and not inn_ok("7707083894")
    assert snils_ok("112-233-445 95") and not snils_ok("112-233-445 96")
    assert luhn_ok("4111 1111 1111 1111") and not luhn_ok("4111 1111 1111 1112")


@pytest.mark.parametrize("text,expected", [
    ("Иванов Иван Иванович", ("Иванов Иван Иванович", "fio")),
    ("ИВАНОВ ИВАН ИВАНОВИЧ", ("ИВАНОВ ИВАН ИВАНОВИЧ", "fio")),
    ("Иван Иванович Иванов", ("Иван Иванович Иванов", "fio")),
    ("подпись Петров А.Б.", ("Петров А.Б.", "fio")),
    ("А.Б. Петров", ("А.Б. Петров", "fio")),
    ("+7 (916) 123-45-67", ("+7 (916) 123-45-67", "phone")),
    ("8-916-123-45-67", ("8-916-123-45-67", "phone")),
    ("ivan.petrov@mail.ru", ("ivan.petrov@mail.ru", "email")),
    ("СНИЛС 112-233-445 95", ("112-233-445 95", "snils")),
    ("ИНН 500100732259", ("500100732259", "inn")),
    ("карта 4111 1111 1111 1111", ("4111 1111 1111 1111", "card")),
    ("паспорт 45 06 123456", ("45 06 123456", "passport")),
    ("серия 4506 № 123456", ("4506 № 123456", "passport")),
    ("дата рождения: 01.02.1990", ("01.02.1990", "birthdate")),
])
def test_detect(text, expected):
    assert expected in kinds(text)


@pytest.mark.parametrize("text", [
    "Договор № 123456789012 от 01.02.2024",
    "Москва Красная Площадь",
    "Сумма 1234567890123",
    "Сумма 4111 1111 1111 1112",
])
def test_no_false_positives(text):
    assert kinds(text) == []


def test_modes():
    text = "Иванов Иван Иванович, тел. +7 916 123-45-67; Иванов Иван Иванович"
    assert Masker("tag").mask_text(text) == "[ФИО], тел. [ТЕЛЕФОН]; [ФИО]"
    assert Masker("pseudo").mask_text(text) == "[ФИО_1], тел. [ТЕЛЕФОН_1]; [ФИО_1]"
    assert "Иван" not in Masker("stars").mask_text(text)


def test_kind_filter():
    m = Masker("tag", {"email"})
    assert m.mask_text("Иванов Иван Иванович a@b.ru") == "Иванов Иван Иванович [EMAIL]"


def _docx_bytes():
    doc = Document()
    p = doc.add_paragraph()
    # ФИО разбито на несколько runs с разным форматированием
    p.add_run("Клиент: Иванов ").bold = True
    p.add_run("Иван ")
    p.add_run("Иванович").italic = True
    p.add_run(", email ivan@mail.ru.")
    t = doc.add_table(rows=1, cols=2)
    t.cell(0, 0).text = "Телефон"
    t.cell(0, 1).text = "+7 916 123-45-67"
    doc.sections[0].header.paragraphs[0].text = "Исполнитель: Петров А.Б."
    doc.core_properties.author = "Секретный Автор"
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def test_docx():
    masker = Masker("tag")
    res = Document(io.BytesIO(mask_docx(_docx_bytes(), masker)))
    body = "\n".join(p.text for p in res.paragraphs)
    assert body == "Клиент: [ФИО], email [EMAIL]."
    assert res.tables[0].cell(0, 1).text == "[ТЕЛЕФОН]"
    assert res.sections[0].header.paragraphs[0].text == "Исполнитель: [ФИО]"
    assert res.core_properties.author == ""
    assert masker.stats == {"fio": 2, "email": 1, "phone": 1}
    assert res.paragraphs[0].runs[0].bold  # форматирование сохранено


def _xlsx_bytes():
    wb = Workbook()
    ws = wb.active
    ws["A1"] = "Иванов Иван Иванович"
    ws["B1"] = 500100732259
    ws["C1"] = 42
    ws["D1"] = '="Иванов Иван Иванович"'
    wb.create_sheet("Лист2")["A1"] = "почта a@b.ru"
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def test_xlsx():
    wb = load_workbook(io.BytesIO(mask_xlsx(_xlsx_bytes(), Masker("tag"))))
    ws = wb.worksheets[0]
    assert ws["A1"].value == "[ФИО]"
    assert ws["B1"].value == "[ИНН]"
    assert ws["C1"].value == 42
    assert ws["D1"].value.startswith("=")
    assert wb.worksheets[1]["A1"].value == "почта [EMAIL]"


def test_api():
    r = client.post(
        "/api/mask",
        files={"file": ("отчёт.docx", _docx_bytes())},
        data={"mode": "pseudo", "kinds": ["fio", "email"]},
    )
    assert r.status_code == 200
    assert "UTF-8''" in r.headers["content-disposition"]
    text = "\n".join(p.text for p in Document(io.BytesIO(r.content)).paragraphs)
    assert "[ФИО_1]" in text and "[EMAIL_1]" in text
    assert client.get("/").status_code == 200
    assert client.get("/api/health").json()["status"] == "ok"


def test_api_errors():
    assert client.post("/api/mask", files={"file": ("a.doc", b"x")}).status_code == 415
    assert client.post("/api/mask", files={"file": ("a.docx", b"not a zip")}).status_code == 422
    assert client.post("/api/mask", files={"file": ("a.docx", b"x")}, data={"mode": "zzz"}).status_code == 400

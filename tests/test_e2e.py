"""End-to-end: загрузка docx/xlsx через API, маскирование и демаскирование."""
import io

from docx import Document
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from masking.app import app

client = TestClient(app)


def _make_docx() -> bytes:
    d = Document()
    p = d.add_paragraph()
    # ФИО разбито на несколько runs с разным форматированием.
    p.add_run("Договор с ").bold = False
    p.add_run("Ивановым ").bold = True
    p.add_run("Иваном ")
    p.add_run("Ивановичем").italic = True
    p.add_run(", г. Москва, ул. Ленина, д. 15, кв. 42, тел +7 916 123-45-67.")
    d.add_paragraph("Передать Иванову Ивану Ивановичу копию, ИНН 500100732259.")
    buf = io.BytesIO(); d.save(buf)
    return buf.getvalue()


def _project(name="Тест") -> int:
    r = client.post("/api/projects", json={"name": name})
    assert r.status_code == 201, r.text
    return r.json()["id"]


def test_health():
    r = client.get("/api/health")
    assert r.status_code == 200
    assert r.json()["status"] == "ok"


def test_mask_unmask_roundtrip_docx():
    pid = _project()
    data = _make_docx()
    r = client.post(f"/api/projects/{pid}/mask", files={"file": ("in.docx", data)})
    assert r.status_code == 200
    masked = r.content
    doc = Document(io.BytesIO(masked))
    text = "\n".join(p.text for p in doc.paragraphs)
    # ФИО маскируется, адрес не маскируется (по умолчанию выключен), а номер
    # дома внутри адреса — да, по правилу DIGITS. Так пропадает возможность
    # восстановить ФИО по номеру дома в соседнем контексте.
    assert "[ФИО_1" in text
    assert "Иванов" not in text
    # ФИО в разных падежах → один человек, один seq.
    assert text.count("[ФИО_1") == 2

    r2 = client.post(f"/api/projects/{pid}/unmask", files={"file": ("masked.docx", masked)})
    assert r2.status_code == 200
    rest = Document(io.BytesIO(r2.content))
    rtext = "\n".join(p.text for p in rest.paragraphs)
    assert "Ивановым Иваном Ивановичем" in rtext
    assert "Иванову Ивану Ивановичу" in rtext


def test_entities_and_history():
    pid = _project("Истор")
    client.post(f"/api/projects/{pid}/mask", files={"file": ("x.docx", _make_docx())})

    ents = client.get(f"/api/projects/{pid}/entities").json()
    kinds = {e["kind"] for e in ents}
    assert kinds >= {"FIO", "PHONE", "INN", "DIGITS"}  # д. 15, кв. 42 → DIGITS
    fio = next(e for e in ents if e["kind"] == "FIO")
    assert fio["canonical"].startswith("Иванов")
    assert fio["occurrences"] == 2

    hist = client.get(f"/api/projects/{pid}/history/mask").json()
    assert len(hist) == 1
    assert hist[0]["output"] is not None
    assert hist[0]["stats"]["total"] >= 3


def test_xlsx_roundtrip():
    pid = _project("Xls")
    wb = Workbook()
    ws = wb.active
    ws["A1"] = "Иванов Иван Иванович"
    ws["B1"] = "г. Москва, ул. Ленина, д. 5"
    ws["C1"] = 42
    ws["D1"] = "=A1"  # формула — не трогаем
    buf = io.BytesIO(); wb.save(buf)

    r = client.post(f"/api/projects/{pid}/mask", files={"file": ("x.xlsx", buf.getvalue())})
    assert r.status_code == 200
    masked = load_workbook(io.BytesIO(r.content)).active
    assert masked["A1"].value.startswith("[ФИО_")
    # B1: адрес не маскируется, но «д. 5» (5 — одна цифра) не трогаем, а
    # потому число останется. Проверяем, что строка не содержит ФИО.
    assert "Иванов" not in str(masked["B1"].value)
    # Число 42 (одна группа ≥ 2 цифр) превращается в токен цифровой маски,
    # хранящийся как строка.
    assert str(masked["C1"].value).startswith("[ЦИФ_")
    assert str(masked["D1"].value).startswith("=")

    r2 = client.post(f"/api/projects/{pid}/unmask", files={"file": ("m.xlsx", r.content)})
    assert r2.status_code == 200
    restored = load_workbook(io.BytesIO(r2.content)).active
    assert restored["A1"].value == "Иванов Иван Иванович"
    assert str(restored["C1"].value) == "42"


def test_unknown_tokens_report():
    pid = _project("Unknown")
    d = Document(); d.add_paragraph("Передать [ФИО_99] и [АДРЕС_42].")
    buf = io.BytesIO(); d.save(buf)
    r = client.post(f"/api/projects/{pid}/unmask", files={"file": ("m.docx", buf.getvalue())})
    assert r.status_code == 200
    import json
    stats = json.loads(r.headers["x-unmasking-stats"])
    assert stats["restored"] == 0
    assert set(stats["unknown"]) == {"[ФИО_99]", "[АДРЕС_42]"}


def test_api_errors():
    pid = _project("Err")
    # неверный формат
    r = client.post(f"/api/projects/{pid}/mask", files={"file": ("a.doc", b"x")})
    assert r.status_code == 415
    # битый docx
    r = client.post(f"/api/projects/{pid}/mask", files={"file": ("a.docx", b"not a zip")})
    assert r.status_code == 422
    # несуществующий проект
    r = client.post("/api/projects/9999/mask", files={"file": ("a.docx", b"x")})
    assert r.status_code == 404


def test_project_delete():
    pid = _project("Удалим")
    client.post(f"/api/projects/{pid}/mask", files={"file": ("in.docx", _make_docx())})
    r = client.delete(f"/api/projects/{pid}")
    assert r.status_code == 204
    assert client.get(f"/api/projects/{pid}").status_code == 404

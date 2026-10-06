"""Маскирование .xlsx: значения ячеек, комментарии, метаданные. Формулы не трогаем."""
import io

from openpyxl import load_workbook

from .masker import Masker


def mask_xlsx(data: bytes, masker: Masker) -> bytes:
    wb = load_workbook(io.BytesIO(data))
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                value = cell.value
                if cell.data_type == "f":
                    continue
                if isinstance(value, bool):
                    pass
                elif isinstance(value, int) or (isinstance(value, float) and value.is_integer()):
                    masked = masker.mask_text(str(int(value)))
                    if masked != str(int(value)):
                        cell.value = masked
                        cell.number_format = "@"
                elif isinstance(value, str) and value:
                    masked = masker.mask_text(value)
                    if masked != value:
                        cell.value = masked
                if cell.comment is not None:
                    text = cell.comment.text
                    masked = masker.mask_text(text)
                    if masked != text:
                        cell.comment.text = masked

    props = wb.properties
    props.creator = ""
    props.lastModifiedBy = ""
    props.title = props.subject = props.description = props.keywords = None

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
